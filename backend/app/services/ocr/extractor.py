from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def extract_pdf_pages(file_path: str) -> list[tuple[str, float, int]]:
    doc = fitz.open(file_path)
    pages: list[tuple[str, float, int]] = []
    for page in doc:
        # 텍스트 블록 수 (get_text("blocks") 는 텍스트 블록만 안정적으로 보고)
        text_block_count = sum(1 for b in page.get_text("blocks") if b[6] == 0)

        # 이미지 면적: get_text("blocks") 는 전면 이미지를 누락하므로 get_image_info 로 직접 집계
        page_area = page.rect.width * page.rect.height
        image_area = sum(
            abs((b["bbox"][2] - b["bbox"][0]) * (b["bbox"][3] - b["bbox"][1]))
            for b in page.get_image_info()
        )
        image_ratio = min(image_area / page_area, 1.0) if page_area > 0 else 0.0

        pages.append((page.get_text(), image_ratio, text_block_count))
    doc.close()
    return pages


def render_pdf_pages(
    file_path: str, indices: set[int], dpi: int = 150
) -> dict[int, Image.Image]:
    if not indices:
        return {}
    doc = fitz.open(file_path)
    images: dict[int, Image.Image] = {}
    try:
        for i in sorted(indices):
            if 0 <= i < doc.page_count:
                pix = doc[i].get_pixmap(dpi=dpi)
                images[i] = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    finally:
        doc.close()
    return images


# 그룹 도형을 재귀적으로 순회하며 (텍스트 줄, 텍스트 블록 수, 이미지 면적) 집계
def _collect_shape_content(shapes) -> tuple[list[str], int, int]:
    lines: list[str] = []
    text_block_count = 0
    image_area = 0

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            g_lines, g_count, g_image_area = _collect_shape_content(shape.shapes)
            lines.extend(g_lines)
            text_block_count += g_count
            # 그룹 자식은 좌표계가 달라 슬라이드 면적 환산이 부정확하므로,
            # 이미지가 있으면 그룹 전체 면적으로 근사 (빈 페이지 오판 방지)
            if g_image_area > 0:
                image_area += shape.width * shape.height
            continue

        if shape.has_text_frame:
            text_block_count += 1
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if line:
                    lines.append(line)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image_area += shape.width * shape.height

    return lines, text_block_count, image_area


def extract_pptx_native_text(file_path: str) -> list[tuple[str, float, int]]:
    prs = Presentation(file_path)
    slide_area = prs.slide_width * prs.slide_height
    results: list[tuple[str, float, int]] = []

    for slide in prs.slides:
        lines, text_block_count, image_area = _collect_shape_content(slide.shapes)
        image_ratio = min(image_area / slide_area, 1.0) if slide_area > 0 else 0.0
        results.append(("\n".join(lines), image_ratio, text_block_count))

    return results


def render_pptx_pages(
    file_path: str, indices: set[int] | None = None, dpi: int = 150
) -> dict[int, Image.Image]:
    """LibreOffice 로 PDF 변환 후 지정한 페이지 인덱스만 이미지로 렌더링.

    indices=None 이면 전 페이지. (LibreOffice 변환 자체는 문서 단위라 불가피하지만,
    무거운 픽스맵 렌더링은 OCR 필요 페이지로 한정한다.)
    """
    path = Path(file_path)
    with tempfile.TemporaryDirectory() as tmpdir:
        try:
            subprocess.run(
                [
                    "libreoffice", "--headless",
                    "--convert-to", "pdf",
                    "--outdir", tmpdir,
                    str(path),
                ],
                check=True,
                timeout=60,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
        except FileNotFoundError as e:
            raise RuntimeError(
                "libreoffice 를 찾을 수 없습니다. PPTX OCR fallback을 위해 LibreOffice를 설치하세요."
            ) from e
        except subprocess.CalledProcessError as e:
            raise RuntimeError(f"LibreOffice PDF 변환 실패: {e}") from e

        pdf_path = Path(tmpdir) / (path.stem + ".pdf")
        doc = fitz.open(str(pdf_path))
        images: dict[int, Image.Image] = {}
        try:
            for i, page in enumerate(doc):
                if indices is not None and i not in indices:
                    continue
                pix = page.get_pixmap(dpi=dpi)
                images[i] = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        finally:
            doc.close()
    return images
